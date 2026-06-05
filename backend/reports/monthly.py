import json
import os
import re
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from lxml import etree
from docx.oxml.ns import qn

REPORTS_DIR = Path("reports")
TEMPLATES_DIR = Path("templates")
STYLE_PATH = Path("REPORT_STYLE.md")
REPORTS_DIR.mkdir(exist_ok=True)

_TYPE_LABEL = {"text": "Note", "email": "Email", "file": "File"}

SECTIONS = {
    "key_projects":       "Key Projects — ongoing projects, deliverables, milestones, operational topics",
    "what_was_good":      "What was good — positive outcomes, above-budget results, achievements, wins",
    "market_competition": "Market and Competition — market trends, competitor activity, pricing intelligence, external factors",
    "less_satisfactory":  "What was less satisfactory — issues, delays, below-expectations results, unresolved challenges",
    "news":               "News — announcements, new partnerships, upcoming events, structural or organisational changes",
}


def find_template() -> Path | None:
    for ext in (".docx", ".html", ".md"):
        p = TEMPLATES_DIR / f"monthly_report_template{ext}"
        if p.exists():
            return p
    return None


def _load_style() -> str:
    return STYLE_PATH.read_text(encoding="utf-8") if STYLE_PATH.exists() else ""


def _normalize_paragraphs(text: str) -> str:
    """Promote single newlines to double so every topic gets a blank line."""
    text = text.strip()
    text = re.sub(r"\n{3,}", "\n\n", text)        # collapse 3+ → 2
    text = re.sub(r"(?<!\n)\n(?!\n)", "\n\n", text)  # single → double
    return text


def _categorize_entries(entries: list) -> dict[str, str]:
    """
    One AI call: categorize all entries into the 5 report sections.
    Returns a dict keyed by section name with paragraph-formatted text.
    """
    empty = {k: "(no entries)" for k in SECTIONS}

    if not entries:
        return empty

    raw = "\n\n".join(
        f"[{_TYPE_LABEL.get(e.type, e.type)}] {e.source}:\n{e.content[:400]}"
        for e in entries
    )

    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        # Fallback: put raw content in key_projects only
        fallback = "\n\n".join(e.content[:300] for e in entries)
        return {**empty, "key_projects": fallback}

    style = _load_style()
    section_list = "\n".join(f"  - {k}: {v}" for k, v in SECTIONS.items())

    prompt = (
        "You are writing an executive monthly business report.\n"
        "Categorize and summarize the following entries into exactly these 5 sections:\n"
        f"{section_list}\n\n"
        "Rules:\n"
        "- EACH distinct topic or project MUST be a separate string in a JSON array\n"
        "- Each string: 1-2 sentences, executive language, no bullet points\n"
        "- No person names. Keep company/customer names and figures.\n"
        "- An entry may contribute to more than one section if relevant.\n"
        "- If no entries fit a section, use the array: [\"(no entries)\"]\n"
        f"Style guide:\n{style}\n\n"
        "Example of correct format:\n"
        '{"key_projects":["First topic sentence.","Second topic sentence.","Third topic."],'
        '"what_was_good":["One good outcome here."],'
        '"market_competition":["Market topic here."],'
        '"less_satisfactory":["Issue here."],'
        '"news":["News item here."]}\n\n'
        "Return ONLY valid JSON where every section value is an array of strings.\n\n"
        f"Entries:\n{raw}"
    )

    try:
        import anthropic
        client = anthropic.Anthropic(api_key=api_key)
        msg = client.messages.create(
            model="claude-haiku-4-5-20251001",
            max_tokens=2000,
            messages=[{"role": "user", "content": prompt}],
        )
        text = msg.content[0].text.strip()

        # Strip markdown code fences if present
        text = re.sub(r"^```[a-z]*\n?", "", text)
        text = re.sub(r"\n?```$", "", text)
        text = text.strip()

        data = json.loads(text)
        result = {}
        for k in SECTIONS:
            val = data.get(k, ["(no entries)"])
            if isinstance(val, list):
                joined = "\n\n".join(s.strip() for s in val if s.strip())
                result[k] = joined or "(no entries)"
            else:
                result[k] = _normalize_paragraphs(str(val))
        return result

    except Exception:
        return {**empty, "key_projects": raw[:1000]}


# ── DOCX multi-paragraph cell writer ──────────────────────────────────────


def _set_cell_paragraphs(cell, text: str) -> None:
    """
    Write content into a cell as proper Word paragraphs.
    - Removes ALL pre-existing paragraphs (clears trailing empties).
    - Inserts one slim blank-line spacer between topics.
    """
    parts = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not parts:
        parts = [""]

    tc = cell._tc
    existing_ps = tc.findall(qn("w:p"))
    if not existing_ps:
        return

    # Capture formatting from the template paragraph before removing anything
    first_p = existing_ps[0]
    pPr_el = first_p.find(qn("w:pPr"))
    rPr_el = None
    for r_el in first_p.findall(qn("w:r")):
        rPr_el = r_el.find(qn("w:rPr"))
        if rPr_el is not None:
            break

    # Remove ALL existing paragraphs from the cell
    for p in existing_ps:
        tc.remove(p)

    def _make_content_para(content: str):
        p = etree.Element(qn("w:p"))
        if pPr_el is not None:
            p.append(deepcopy(pPr_el))
        r = etree.SubElement(p, qn("w:r"))
        if rPr_el is not None:
            r.append(deepcopy(rPr_el))
        t = etree.SubElement(r, qn("w:t"))
        t.text = content
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        return p

    def _make_spacer_para():
        """One blank line between topics — single line height, no extra before/after."""
        p = etree.Element(qn("w:p"))
        pPr = etree.SubElement(p, qn("w:pPr"))
        spacing = etree.SubElement(pPr, qn("w:spacing"))
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "240")       # 12pt — one normal line
        spacing.set(qn("w:lineRule"), "auto")
        return p

    # Append paragraphs directly into the tc element
    # Word requires at least one <w:p> in every <w:tc>
    for i, part in enumerate(parts):
        if i > 0:
            tc.append(_make_spacer_para())
        tc.append(_make_content_para(part))


def _replace_docx(doc, replacements: dict[str, str]) -> None:
    """Replace all placeholder tokens in paragraphs and table cells."""

    def _apply_para(para):
        full = para.text
        new_full = full
        for ph, val in replacements.items():
            new_full = new_full.replace(ph, val)
        if new_full != full:
            if para.runs:
                para.runs[0].text = new_full
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new_full)

    for para in doc.paragraphs:
        _apply_para(para)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for ph, val in replacements.items():
                    if ph in cell.text:
                        _set_cell_paragraphs(cell, val)
                        break  # one placeholder per cell

    # Zero out body-level empty paragraphs between/after tables so they
    # don't push content onto a second page.
    for child in doc.element.body:
        if child.tag == qn("w:p") and not child.text:
            pPr = child.find(qn("w:pPr"))
            if pPr is None:
                pPr = etree.SubElement(child, qn("w:pPr"))
            spacing = pPr.find(qn("w:spacing"))
            if spacing is None:
                spacing = etree.SubElement(pPr, qn("w:spacing"))
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")
            spacing.set(qn("w:line"), "1")
            spacing.set(qn("w:lineRule"), "exact")


# ── Main generator ─────────────────────────────────────────────────────────


def generate_monthly_report(entries: list, year: int, month: int) -> Path:
    template_path = find_template()
    if not template_path:
        raise FileNotFoundError(
            "No monthly report template found in templates/. "
            "Please upload a template (.md, .html, or .docx) first."
        )

    now = datetime.now(timezone.utc)
    month_str = datetime(year, month, 1).strftime("%B %Y")

    filtered = [
        e for e in entries
        if e.date_stamp.year == year and e.date_stamp.month == month
    ]

    categories = _categorize_entries(filtered)

    replacements = {
        "{{month}}": month_str,
        # Current section tokens
        "{{key_projects}}":       categories["key_projects"],
        "{{what_was_good}}":      categories["what_was_good"],
        "{{market_competition}}": categories["market_competition"],
        "{{less_satisfactory}}":  categories["less_satisfactory"],
        "{{news}}":               categories["news"],
        # Legacy tokens for backward compatibility
        "{{entries}}":            categories["key_projects"],
        "{{summary}}":            categories["what_was_good"],
        "{{topic_summary}}":      categories["market_competition"],
    }

    ext = template_path.suffix.lower()
    ts = now.strftime("%Y%m%dT%H%M")

    if ext in (".md", ".html"):
        content = template_path.read_text(encoding="utf-8")
        for ph, val in replacements.items():
            content = content.replace(ph, val)
        filename = f"monthly_{year}-{month:02d}_generated-{ts}{ext}"
        output_path = REPORTS_DIR / filename
        output_path.write_text(content, encoding="utf-8")
        return output_path

    if ext == ".docx":
        from docx import Document
        doc = Document(str(template_path))
        _replace_docx(doc, replacements)
        filename = f"monthly_{year}-{month:02d}_generated-{ts}.docx"
        output_path = REPORTS_DIR / filename
        doc.save(str(output_path))
        return output_path

    raise ValueError(f"Unsupported template format: {ext}")
